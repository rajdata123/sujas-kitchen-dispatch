"""
Generates the per-branch "Food Dispatch Time & Temperature Log" workbook —
same 2-tab format (Catergery-1 / Catergery-2) Shapan currently builds by
hand, grouped into prep/storage-zone buckets with a yellow banner row per
bucket, S.No restarting at 1 within each bucket, and blank Departure/Arrival
Temperature columns left for the driver to fill in during real dispatch.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from item_master import SHEET_GROUP_ORDER

COLS = ["S.no", "Item Name", "Category", "UOM", "Quantity",
        "Departure Temp. (⁰C)", "Arrival Temp. (⁰C)"]
COL_WIDTHS = [5.14, 36.14, 13, 11, 10, 12, 12]

YELLOW = PatternFill("solid", fgColor="FFFFFF00")
THIN = Side(style="thin")
MEDIUM = Side(style="medium")
BOX_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX_MED = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)


def _write_header_block(ws, branch: str, delivery_date: str):
    ws.merge_cells("A1:G1")
    ws["A1"] = "FOOD DISPATCH TIME & TEMPERATURE LOG"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].border = BOX_THIN
    ws.row_dimensions[1].height = 40

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Order No. & Location:  {branch}" + " " * 60 + f"Date: {delivery_date}"
    ws.merge_cells("A3:G3")
    ws["A3"] = "Meal Type:" + " " * 45 + "Checked By:" + " " * 45 + "Dispatch Time: "
    ws.merge_cells("A4:G4")
    ws["A4"] = "Driver Name:" + " " * 42 + "Vehicle No:" + " " * 46 + "Delivery Time: "
    for row in (2, 3, 4):
        cell = ws.cell(row=row, column=1)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX_THIN
        ws.row_dimensions[row].height = 18.75
    ws.row_dimensions[5].height = 10.5


def _write_column_headers(ws, row: int):
    for c, label in enumerate(COLS, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = Font(name="Arial", size=9, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 24


def _write_group_banner(ws, row: int, group_name: str):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = ws.cell(row=row, column=1, value=group_name)
    cell.fill = YELLOW
    cell.font = Font(name="Times New Roman", size=11, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 15.75


def _write_item_row(ws, row: int, s_no: int, item_name: str, item_code: str,
                     storage_category: str, uom: str, qty: float):
    values = [s_no, f"{item_name}\n{item_code}", storage_category, uom, qty, None, None]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(name="Arial", size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX_MED
    ws.row_dimensions[row].height = 23.25


def _build_sheet(ws, branch: str, delivery_date: str, groups: dict, group_order: list):
    _write_header_block(ws, branch, delivery_date)
    row = 6
    for group_name in group_order:
        items = groups.get(group_name)
        if not items:
            continue
        _write_column_headers(ws, row)
        row += 1
        _write_group_banner(ws, row, group_name)
        row += 1
        for i, item in enumerate(items, start=1):
            _write_item_row(ws, row, i, item["item_name"], item["item_code"],
                             item["storage_category"], item["uom"], item["qty"])
            row += 1
    for c, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width


def build_branch_workbook(branch: str, delivery_date: str, items: list) -> Workbook:
    """items: list of dicts each with item_name, item_code, uom, qty,
    dispatch_sheet, group, storage_category (already joined against the Item
    Master). Items with a blank dispatch_sheet/group (not yet tagged) are
    skipped here — the caller should have already blocked on unmapped items
    or excluded them explicitly."""
    by_sheet = {"Catergery-1": {}, "Catergery-2": {}}
    for item in items:
        sheet = item.get("dispatch_sheet")
        group = item.get("group")
        if sheet not in by_sheet or not group:
            continue
        by_sheet[sheet].setdefault(group, []).append(item)

    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in ["Catergery-1", "Catergery-2"]:
        ws = wb.create_sheet(sheet_name)
        # Known buckets (from SHEET_GROUP_ORDER) print first, in their usual
        # order; a bucket Shapan created on the fly while tagging a new item
        # (via "+ New group...") isn't in that fixed list yet, so it's
        # appended at the end instead of silently dropped.
        known_order = SHEET_GROUP_ORDER[sheet_name]
        extra_groups = [g for g in by_sheet[sheet_name] if g not in known_order]
        group_order = known_order + extra_groups
        _build_sheet(ws, branch, delivery_date, by_sheet[sheet_name], group_order)
    return wb


def safe_filename(branch: str, delivery_date: str) -> str:
    branch_part = branch.replace(" ", "_")
    date_part = delivery_date.replace("/", "_")
    return f"{branch_part}_Delivery_{date_part}.xlsx"
